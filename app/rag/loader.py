"""文档加载：支持 txt / md / pdf / docx。"""
from pathlib import Path

import pypdf
from docx import Document


def load_document(path: Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        reader = pypdf.PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # 表格内容按行提取，单元格以 | 连接，避免表格数据丢失
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
        return "\n".join(parts).strip()
    raise ValueError(f"暂不支持的文件类型: {suffix}")